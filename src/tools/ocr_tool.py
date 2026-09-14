"""
OCR & Document Parser Tool for the Sovereign Agentic AI Workbench.
Extracts structured piping inspection measurements, thresholds, and metadata.
"""

import os
import re
import json
import base64
from typing import Dict, Any, List
import requests  # Or your local LiteLLM / Ollama client


def encode_image_base64(filepath: str) -> str:
    with open(filepath, "rb") as f:
        return base64.b64encode(f.read()).decode("utf-8")


def parse_image_with_vision_model(filepath: str) -> str:
    """Passes the image to the local Vision LLM (e.g., Qwen2.5-VL via local Ollama/vLLM)."""
    base64_image = encode_image_base64(filepath)

    payload = {
        "model": "qwen2.5-vl",  # Local vision model
        "messages": [
            {
                "role": "user",
                "content": [
                    {
                        "type": "text",
                        "text": (
                            "Extract all text, tables, and diagram descriptions from this image. "
                            "If this is an inspection report, extract all metadata and measurement points."
                        ),
                    },
                    {
                        "type": "image_url",
                        "image_url": {"url": f"data:image/png;base64,{base64_image}"},
                    },
                ],
            }
        ],
        "temperature": 0.0,
    }

    try:
        # Calls local air-gapped inference endpoint
        response = requests.post(
            "http://localhost:11434/v1/chat/completions",
            json=payload,
            timeout=30,
        )
        response.raise_for_status()
        return response.json()["choices"][0]["message"]["content"]
    except Exception as e:
        raise RuntimeError(f"Vision extraction failed: {e}")


def extract_text_from_excel(filepath: str) -> str:
    """Extract cell values and table structure from an Excel workbook (.xlsx / .xls)."""
    chunks = []
    try:
        import openpyxl
        wb = openpyxl.load_workbook(filepath, data_only=True)
        for sheetname in wb.sheetnames:
            sheet = wb[sheetname]
            chunks.append(f"--- Sheet: {sheetname} ---")
            for row in sheet.iter_rows(values_only=True):
                if any(cell is not None for cell in row):
                    row_str = " | ".join(str(cell).strip() if cell is not None else "" for cell in row)
                    chunks.append(row_str)
        return "\n".join(chunks)
    except Exception:
        pass

    try:
        import pandas as pd
        excel_file = pd.ExcelFile(filepath)
        for sheet_name in excel_file.sheet_names:
            df = pd.read_excel(excel_file, sheet_name=sheet_name)
            chunks.append(f"--- Sheet: {sheet_name} ---")
            chunks.append(df.to_string(index=False))
        return "\n".join(chunks)
    except Exception as e:
        return f"[Error extracting Excel content: {e}]"


def validate_inspection_relevance(content: str) -> tuple:
    """
    Validate whether extracted text contains valid industrial inspection telemetry.
    Returns: (is_valid: bool, document_type: str, reason: str)
    """
    if not content or not content.strip():
        return False, "EMPTY_DOCUMENT", "Uploaded document is empty or unreadable."

    # Prevent raw ZIP or binary headers from passing
    if content.startswith("PK\x03\x04") or "Content_Types].xml" in content:
        return False, "CORRUPTED_BINARY", "Uploaded document could not be decoded and contains raw archive headers."

    text_lower = content.lower()

    # Domain keywords indicating industrial piping / NDT / inspection
    inspection_keywords = [
        "inspection", "thickness", "ultrasonic", "api 570", "asme", "corrosion",
        "t_min", "t-min", "t_thresh", "nominal", "wall loss", "piping", "pipe",
        "spool", "ndt", "nde", "ut-", "cml", "tml", "pressure vessel", "hydrotest",
        "schedule 40", "schedule 80", "carbon steel", "alloy", "remaining life",
        "flange", "weld", "cui", "circuit", "caliper", "probe", "retire"
    ]

    matched_keywords = [kw for kw in inspection_keywords if kw in text_lower]

    # Look for characteristic point tags (e.g. RG-01, T-01, P-02, UT-104) followed by numbers
    has_point_readings = bool(re.search(r"\b[A-Za-z]{1,6}-?\d+\b[^\n\d]*\d+\.\d+", content))

    if len(matched_keywords) == 0 and not has_point_readings:
        # Detect non-inspection categories for clearer user feedback
        if any(w in text_lower for w in ["invoice", "goods", "quantity", "rate", "amount", "total", "mrp", "hsn", "diwali"]):
            doc_type = "COMMERCIAL_INVOICE"
            reason = "Document contains commercial invoice / sales billing items with no NDT inspection telemetry."
        elif any(w in text_lower for w in ["resume", "curriculum vitae", "experience", "education", "skills"]):
            doc_type = "RESUME_CV"
            reason = "Document appears to be a resume/CV rather than an engineering inspection log."
        else:
            doc_type = "IRRELEVANT_NON_INSPECTION"
            reason = "No ultrasonic thickness readings, piping inspection points, or API 570 / ASME parameters detected."
        return False, doc_type, reason

    return True, "INSPECTION_LOG", "Valid industrial inspection telemetry detected."


def parse_inspection_document(filepath: str) -> Dict[str, Any]:
    """Parse a scanned report, image, PDF, Excel, or text file into structured inspection data."""
    if not os.path.isfile(filepath):
        raise FileNotFoundError(f"Inspection file not found: {filepath}")

    ext = os.path.splitext(filepath)[1].lower()
    content = ""

    # Route 1: PDF files
    if ext == ".pdf":
        try:
            import pypdf

            reader = pypdf.PdfReader(filepath)
            for page in reader.pages:
                content += (page.extract_text() or "") + "\n"
        except Exception as e:
            raise RuntimeError(f"Failed to parse PDF document: {e}")

    # Route 2: Image files (PNG, JPG)
    elif ext in [".png", ".jpg", ".jpeg", ".webp"]:
        content = parse_image_with_vision_model(filepath)

    # Route 3: Excel files (XLSX, XLS)
    elif ext in [".xlsx", ".xls"]:
        content = extract_text_from_excel(filepath)

    # Route 4: Plain text / CSV files
    else:
        with open(filepath, "r", encoding="utf-8", errors="replace") as f:
            content = f.read()

    return parse_inspection_text(content)


def parse_inspection_text(content: str) -> Dict[str, Any]:
    """Parse raw text content and extract structured metadata, thresholds, and measurements."""
    is_valid, doc_type, reason = validate_inspection_relevance(content)
    if not is_valid:
        return {
            "raw_text": content,
            "document_type": doc_type,
            "is_valid_document": False,
            "status": "INVALID_DOCUMENT",
            "metadata": {"summary": content[:500]},
            "thresholds": {},
            "measurement_points": [],
            "critical_points": [],
            "warning": "Uploaded file or text does not appear to be an industrial inspection log.",
            "rejection_reason": reason,
        }

    # Extract Key Metadata
    metadata = {
        "report_id": _extract_regex(content, r"(?:Report ID|Report Reference ID|Identifier)\s*[:=]\s*([^\n]+)", "UT-SCAN-UNKNOWN"),
        "facility": _extract_regex(content, r"Facility(?:\s*Name)?\s*[:=]\s*([^\n]+)", "Strategic Industrial Facility"),
        "line_number": _extract_regex(content, r"(?:Line Identifier|Line Number|System Component|Line Tag)\s*[:=]\s*([^\(\n]+)", "Unknown Line"),
        "material": _extract_regex(content, r"(?:Pipe Material|Material Grade|Material)\s*[:=]\s*([^\(\n]+)", "Carbon Steel"),
        "standard": _extract_regex(content, r"(?:Applicable Standards?|Design Code / SOP|Standard)\s*[:=]\s*([^\(\n]+)", "API 570 / ASME B31.3"),
        "inspector": _extract_regex(content, r"(?:Inspector|Lead Inspector)\s*[:=]\s*([^\(\n]+)", "Unassigned"),
    }

    # Extract Regulatory Thresholds
    m_min = re.search(r"(?:Retirement Thickness \(T_min\)|T_min|Structural Retirement Thickness)[^\:\=\n]*[:=]\s*([\d\.]+)", content, re.IGNORECASE)
    m_margin = re.search(r"(?:Corrosion Safety Allowance|Margin)[^\:\=\n]*[:=]\s*([\d\.]+)", content, re.IGNORECASE)
    m_thresh = re.search(r"(?:Critical Retirement Threshold|T_threshold|T_thresh)[^\n:]*:\s*([\d\.]+)", content, re.IGNORECASE)

    t_min = float(m_min.group(1)) if m_min else 2.80
    margin = float(m_margin.group(1)) if m_margin else 0.50
    t_thresh = float(m_thresh.group(1)) if m_thresh else round(t_min + margin, 2)

    thresholds = {
        "t_min": t_min,
        "margin": margin,
        "t_threshold": t_thresh,
    }

    # Extract Tabular Measurement Points
    points: List[Dict[str, Any]] = []
    seen_pids = set()

    # Look for nominal wall thickness in document metadata
    doc_nominal = float(_extract_regex(content, r"(?:Nominal Original Wall Thickness|Nominal Thickness|Original Thickness)\s*[:=]\s*([\d\.]+)", "0.0"))

    for raw_line in content.split("\n"):
        line = raw_line.strip()
        # Matches points like RG-01, T-01, UT-104, P-02, CW-05, etc.
        m_pid = re.match(r"^([A-Za-z]{1,6}-?\d+)\b\s*\|?\s*(.*)", line)
        if not m_pid:
            continue
        pid = m_pid.group(1).upper()
        if pid in seen_pids:
            continue
        rest = m_pid.group(2)
        # Find all decimal floating-point numbers on this row
        num_matches = list(re.finditer(r"\b\d+\.\d+\b", rest))
        if not num_matches:
            continue
        first_num_pos = num_matches[0].start()
        desc = rest[:first_num_pos].replace("|", "").strip()
        # Avoid matching narrative observation sentences that begin with a point tag
        if len(desc) > 65 or "exhibit" in desc.lower() or "wall loss" in desc.lower() or "measured at" in desc.lower():
            continue

        nums = [float(m.group(0)) for m in num_matches]
        if len(nums) == 1:
            measured = nums[0]
            nominal = doc_nominal if doc_nominal > 0 else measured
        elif len(nums) == 2:
            nominal = max(nums)
            measured = min(nums)
        else:
            # Multi-column table (e.g. Nominal, Prev 2026, Measured):
            # First is nominal, last is current measured thickness
            nominal = nums[0] if doc_nominal == 0 else doc_nominal
            measured = nums[-1]

        points.append({
            "point_id": pid,
            "description": desc if desc else "Piping Inspection Point",
            "nominal_mm": nominal,
            "measured_mm": measured,
            "is_breached": measured < thresholds["t_threshold"],
        })
        seen_pids.add(pid)

    return {
        "raw_text": content,
        "document_type": "INSPECTION_LOG",
        "is_valid_document": True,
        "status": "VALID_INSPECTION",
        "metadata": metadata,
        "thresholds": thresholds,
        "measurement_points": points,
        "critical_points": [p for p in points if p["is_breached"]],
    }


def _extract_regex(text: str, pattern: str, default: str) -> str:
    match = re.search(pattern, text, re.IGNORECASE)
    return match.group(1).strip() if match else default